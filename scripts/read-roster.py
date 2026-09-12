"""Read-only XLSX adapter. Emits private data to the local matching process only."""
import sys
import json
import hashlib
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = book["Participants"]
headers = [c.value for c in next(sheet.iter_rows())]
fields = ["timestamp", "formEmail", "name", "email", "membershipNumber", "membershipExpiry",
          "membershipConfirmed", "chapterRaw", "credentialRaw", "hoursRaw", "practiceAreas",
          "timezoneRaw", "availabilityRaw", "languageRaw", "englishAnswer", "hasConstraints",
          "schedulingNotes", "goals", "developmentGoals", "commitment"]
expected = ["Timestamp", "Email Address", "Full Name", "What is your preferred email",
            "Your ICF membership", "Membership expired", "Memebership Confirmation",
            "Which Chapter", "What is your most recent", "How many hours", "What are your main areas",
            "What time zone", "Which days", "What are your language", "Would you be comfortable",
            "Do you have any scheduling", "Please list any other", "What do you wish",
            "What coaching competencies", "Thank you for your commitment"]
for index, prefix in enumerate(expected):
    if not str(headers[index] or "").strip().startswith(prefix):
        raise ValueError(f"Unexpected source header in column {index + 1}")
participants = []
for source_row, cells in enumerate(sheet.iter_rows(min_row=2), 2):
    values = [c.value for c in cells]
    if not any(values[i] for i in (1, 2, 3)):
        continue
    row = dict(zip(fields, values[:20]))
    identity = str(row["email"] or row["formEmail"] or "").strip().lower()
    row["id"] = "P-" + hashlib.sha256(identity.encode()).hexdigest()[:12]
    row["sourceRow"] = source_row
    # Membership numbers and personal-development narratives are not needed for matching.
    for field in ("membershipNumber", "membershipExpiry", "goals", "developmentGoals", "timestamp", "formEmail"):
        row.pop(field)
    participants.append(row)
if len({p["id"] for p in participants}) != len(participants):
    raise ValueError("Duplicate participant identities: review the source before matching")
print(json.dumps({"sourceName": source.name, "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                  "participants": participants}, ensure_ascii=False, default=str))
